"""CLI 冒烟（不依赖 UI）。"""

import json

from PIL import Image

from legoart import cli


def test_generate_cli_end_to_end(tmp_path):
    src = tmp_path / "src.png"
    Image.new("RGB", (60, 30), (150, 40, 40)).save(src)
    out = tmp_path / "plan.json"
    prev = tmp_path / "preview.png"

    rc = cli.main(
        [
            "generate",
            "--image", str(src),
            "--width", "12",
            "--color-set", "recommended",
            "--out", str(out),
            "--preview", str(prev),
        ]
    )
    assert rc == 0
    assert out.exists() and prev.exists()
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["grid"]["w"] == 12
    assert data["grid"]["h"] == 6  # 60:30 比例自动
    im = Image.open(prev)
    assert im.size == (12 * 20, 6 * 20)


def test_generate_cm_unit(tmp_path):
    src = tmp_path / "src.png"
    Image.new("RGB", (80, 40), (0, 90, 150)).save(src)
    out = tmp_path / "plan.json"
    rc = cli.main(
        ["generate", "--image", str(src), "--width", "8", "--unit", "cm", "--out", str(out)]
    )
    assert rc == 0
    data = json.loads(out.read_text(encoding="utf-8"))
    assert data["grid"]["w"] == 10  # 8cm / 0.8cm-per-stud
    assert data["grid"]["h"] == 5


def test_cli_export_excel_pdf(tmp_path):
    from openpyxl import load_workbook

    src = tmp_path / "src.png"
    Image.new("RGB", (80, 40), (10, 90, 160)).save(src)
    xlsx = tmp_path / "bom.xlsx"
    pdf = tmp_path / "doc.pdf"
    rc = cli.main(
        [
            "generate", "--image", str(src), "--width", "12",
            "--no-saliency",
            "--excel", str(xlsx), "--pdf", str(pdf),
        ]
    )
    assert rc == 0
    assert xlsx.exists() and pdf.exists()
    assert {"物料清单", "方案参数"} <= set(load_workbook(xlsx).sheetnames)
    assert pdf.read_bytes()[:5] == b"%PDF-"


def test_cli_error_missing_image(tmp_path, capsys):
    rc = cli.main(
        ["generate", "--image", str(tmp_path / "nope.png"), "--width", "8", "--out", str(tmp_path / "p.json")]
    )
    err = capsys.readouterr().err
    assert rc == 1
    assert "图片" in err or "ERR" in err
