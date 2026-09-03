"""导出测试：Excel BOM 结构 / PDF 生成 / api.export_plan。"""

import json

import pytest
from PIL import Image
from openpyxl import load_workbook

from legoart import api
from legoart.inventory import solve_inventory
from legoart.model import ImageSpec
from legoart.model.image_spec import ShortageStrategy


def _plan(w=16, h=16, saliency=False):
    spec = ImageSpec(grid_w=w, grid_h=h, saliency_enabled=saliency)
    img = Image.new("RGB", (64, 64), (150, 40, 40))
    return api.generate_plan(img, spec)


def test_excel_bom_structure(tmp_path):
    plan = _plan()
    cat = api.load_catalog()
    pal = api.build_palette(cat)
    p = tmp_path / "bom.xlsx"
    api.export_plan(plan, excel_path=str(p), catalog=cat, palette=pal)
    wb = load_workbook(p)
    assert {"方案参数", "物料清单"} <= set(wb.sheetnames)
    ws = wb["物料清单"]
    header = [c.value for c in ws[1]]
    assert header[0] == "设计号" and "总用量" in header and "L0" in header
    total = sum(row[-1] for row in ws.iter_rows(min_row=2, values_only=True))
    assert total == len(plan.placements)
    layer_total = sum(row[header.index("L0")] for row in ws.iter_rows(min_row=2, values_only=True))
    assert layer_total == sum(1 for p in plan.placements if p.layer == 0)


def test_excel_solver_sheets(tmp_path):
    plan = _plan(8, 8)
    cat = api.load_catalog()
    pal = api.build_palette(cat)
    solver = solve_inventory(
        plan.placements, {}, strategy=ShortageStrategy.MISSING_LIST, catalog=cat, palette=pal
    )
    assert solver.missing
    p = tmp_path / "bom2.xlsx"
    api.export_plan(plan, excel_path=str(p), solver=solver, catalog=cat, palette=pal)
    wb = load_workbook(p)
    assert "缺件清单" in wb.sheetnames
    ws = wb["缺件清单"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(solver.missing)


def test_pdf_generation(tmp_path):
    plan = _plan(16, 16)
    cat = api.load_catalog()
    pal = api.build_palette(cat)
    p = tmp_path / "doc.pdf"
    api.export_plan(plan, pdf_path=str(p), catalog=cat, palette=pal)
    data = p.read_bytes()
    assert data[:5] == b"%PDF-"
    assert len(data) > 4000
    assert data.count(b"/Type /Page") >= 2  # 封面+内容页


def test_pdf_with_solver_and_layers(tmp_path):
    from legoart.inventory import solve_inventory

    # 加一层凸起使说明书含多物理层
    import numpy as np

    from legoart.model import Region

    mask = np.zeros((16, 16), dtype=bool)
    mask[4:8, 4:8] = True
    spec = ImageSpec(grid_w=16, grid_h=16, saliency_enabled=False)
    img = Image.new("RGB", (64, 64), (40, 90, 180))
    plan = api.generate_plan(img, spec, regions=[Region(mask=mask, raise_layers=2, source="user")])
    cat = api.load_catalog()
    pal = api.build_palette(cat)
    solver = solve_inventory(
        plan.placements, {}, strategy=ShortageStrategy.MISSING_LIST, catalog=cat, palette=pal
    )
    p = tmp_path / "doc2.pdf"
    api.export_plan(plan, pdf_path=str(p), solver=solver, catalog=cat, palette=pal)
    data = p.read_bytes()
    assert data[:5] == b"%PDF-"
    assert data.count(b"/Type /Page") >= 4  # 封面+清单+缺件+多层页


def test_font_register_returns_name():
    from legoart.export.fonts import register_cn_font

    name = register_cn_font()
    assert name in ("CNFont", "Helvetica")


def test_export_requires_path():
    plan = _plan(4, 4)
    with pytest.raises(ValueError):
        api.export_plan(plan)
