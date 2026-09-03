"""库存 Excel 导入/导出测试。"""

import pytest

from legoart.errors import StorageError
from legoart.storage.inventory_io import export_inventory_xlsx, import_inventory_xlsx


def test_roundtrip(tmp_path):
    p = tmp_path / "inv.xlsx"
    rows = [
        {"design_id": "3024", "color_id": "Red", "quantity": 120, "note": "拆包"},
        {"design_id": "3020", "color_id": "Blue", "quantity": 8, "note": ""},
    ]
    export_inventory_xlsx(p, rows)
    back = import_inventory_xlsx(p)
    assert len(back) == 2
    assert back[0]["design_id"] == "3024" and back[0]["quantity"] == 120
    assert back[1]["color_id"] == "Blue"


def test_import_ignores_nonpositive_and_blank(tmp_path):
    from openpyxl import Workbook

    p = tmp_path / "in.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["design_id", "color_id", "quantity", "note"])
    ws.append(["3024", "Red", 0, "skip"])
    ws.append(["3024", "Red", -5, "skip"])
    ws.append(["3023", "White", 3, None])
    ws.append([None, None, None, None])
    wb.save(str(p))
    rows = import_inventory_xlsx(p)
    assert len(rows) == 1 and rows[0]["design_id"] == "3023" and rows[0]["quantity"] == 3


def test_bad_header(tmp_path):
    from openpyxl import Workbook

    p = tmp_path / "bad.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["a", "b"])
    ws.append([1, 2])
    wb.save(str(p))
    with pytest.raises(StorageError):
        import_inventory_xlsx(p)


def test_bad_quantity(tmp_path):
    from openpyxl import Workbook

    p = tmp_path / "badq.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["design_id", "color_id", "quantity", "note"])
    ws.append(["3024", "Red", "很多", ""])
    wb.save(str(p))
    with pytest.raises(StorageError):
        import_inventory_xlsx(p)


def test_missing_file():
    with pytest.raises(StorageError):
        import_inventory_xlsx("no-such.xlsx")
