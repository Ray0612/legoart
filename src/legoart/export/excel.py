"""Excel BOM 导出（openpyxl）。

Sheet「方案参数」→「物料清单」（设计号/名称/颜色/RGB 色块/逐层用量/总用量）
→「缺件清单」「替代明细」（库存模式求解后出现）。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..catalog import Catalog
from ..color import Palette
from ..model import MosaicPlan

_PARAM_ROWS: list[tuple[str, str]] = []


def _hex(catalog: Catalog, color_id: str) -> str:
    c = catalog.color(color_id)
    return c.rgb_hex if c else "#888888"


def _fill_from_hex(hex_str: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_str.lstrip("#") or "888888")


def build_bom_xlsx(
    path: str | Path,
    plan: MosaicPlan,
    palette: Palette | None = None,  # noqa: ARG001  # 保留签名一致性
    catalog: Catalog | None = None,
    solver=None,
) -> Path:
    from ..api import load_catalog

    cat = catalog or load_catalog()
    spec = plan.spec
    levels = sorted(plan.stack.layers, key=lambda x: x.level) if plan.stack else []
    level_cols = [f"L{ly.level}" for ly in levels]

    # 逐层用量 (design,color)->{layer:qty} 与总用量
    per: dict[tuple[str, str], dict[int, int]] = {}
    totals: dict[tuple[str, str], int] = {}
    for item in plan.bom.items:
        key = (item.design_id, item.color_id)
        totals[key] = totals.get(key, 0) + item.qty
        if item.layer is not None:
            per.setdefault(key, {})[item.layer] = per.get(key, {}).get(item.layer, 0) + item.qty
    rows = sorted(totals.keys())

    wb = Workbook()
    ws_param = wb.active
    ws_param.title = "方案参数"
    params = [
        ("网格(W×H studs)", f"{plan.grid.width} × {plan.grid.height}"),
        ("总格数", str(plan.grid.width * plan.grid.height)),
        ("风格", spec.style.value),
        ("需求轴", "库存约束" if spec.use_inventory else "理想购物车"),
        ("缺货策略", spec.shortage_strategy.value if spec.use_inventory else "-"),
        ("候选色集", spec.color_set),
        ("物理层数", str(plan.stack.height_total if plan.stack else 1)),
        ("合并后板件数", str(len(plan.placements))),
        ("生成耗时 ms", str(plan.elapsed_ms)),
        ("目录版本", spec.catalog_version or "-"),
        ("生成时间", plan.generated_at),
    ]
    for k, v in params:
        ws_param.append([k, v])
    if plan.warnings:
        ws_param.append([])
        ws_param.append(["预警", ""])
        for wmsg in plan.warnings:
            ws_param.append(["", wmsg])
    ws_param.column_dimensions["A"].width = 22
    ws_param.column_dimensions["B"].width = 60

    ws_bom = wb.create_sheet("物料清单")
    header = ["设计号", "零件名称", "颜色名称", "RGB"] + level_cols + ["总用量"]
    ws_bom.append(header)
    for c in ws_bom[1]:
        c.font = Font(bold=True)
    for key in rows:
        design_id, color_id = key
        part = cat.part(design_id)
        col_spec = cat.color(color_id)
        color_name = col_spec.name_bl if col_spec else color_id
        hex_str = _hex(cat, color_id)
        line = [
            design_id,
            part.name if part else design_id,
            color_name,
            hex_str,
        ]
        for ly in levels:
            line.append(per.get(key, {}).get(ly.level, 0))
        line.append(totals[key])
        ws_bom.append(line)
        r_idx = ws_bom.max_row
        ws_bom.cell(r_idx, 4).fill = _fill_from_hex(hex_str)
        ws_bom.cell(r_idx, 4).alignment = Alignment(horizontal="center")
    widths = [10, 26, 16, 12] + [7] * len(level_cols) + [9]
    for i, wdt in enumerate(widths, start=1):
        ws_bom.column_dimensions[get_column_letter(i)].width = wdt
    ws_bom.freeze_panes = "A2"

    if solver is not None:
        if getattr(solver, "missing", None):
            ws_m = wb.create_sheet("缺件清单")
            ws_m.append(["设计号", "颜色", "数量"])
            for m in solver.missing:
                ws_m.append([m.design_id, m.color_id, m.qty])
            for c in ws_m[1]:
                c.font = Font(bold=True)
            for col, wdt in zip("ABC", (10, 20, 8), strict=True):
                ws_m.column_dimensions[col].width = wdt
        if getattr(solver, "substituted", None):
            ws_s = wb.create_sheet("替代明细")
            ws_s.append(["原设计号", "原颜色", "替换设计号", "替换颜色", "件数(每处)"])
            for orig, d, k, c in solver.substituted:
                ws_s.append([orig.design_id, orig.color_id, d, c, k])
            for c in ws_s[1]:
                c.font = Font(bold=True)
            for col, wdt in zip("ABCDE", (10, 16, 10, 16, 12), strict=True):
                ws_s.column_dimensions[col].width = wdt

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return out
